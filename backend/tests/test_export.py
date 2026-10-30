"""Round-trip tests for full-account export and filtered transaction export."""

from __future__ import annotations

import csv
import io
import json
import sys
from datetime import UTC, datetime
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from main import app
from routers.upload import get_action_state
from sdk.export_bundle import DATASET_NAMES, DATASETS, BackupValidationError, parse_backup

EXPECTED_COLUMNS = {
    "transactions": (
        "id",
        "date",
        "description",
        "amount",
        "category",
        "source_file",
        "category_edited",
    ),
    "budgets": ("category", "monthly_budget"),
    "category_rules": ("category", "keywords"),
    "subscriptions": (
        "stream_id",
        "merchant",
        "dominant_category",
        "cadence",
        "amount",
        "baseline_amount",
        "expected_amount",
        "last_paid_amount",
        "confidence",
        "active",
        "ignored",
        "essential",
        "manually_managed",
        "price_increase",
        "cancellation_candidate",
        "negotiation_opportunity",
        "is_new_recurring",
        "missed_expected_charge",
        "charge_count",
        "trend",
        "last_charge_date",
        "next_expected_charge_date",
        "next_due_date",
        "status_group",
        "payment_state",
    ),
    "subscription_charges": ("stream_id", "date", "amount"),
    "goals": (
        "id",
        "name",
        "target_amount",
        "target_date",
        "priority",
        "category",
        "status",
        "created_at",
        "updated_at",
        "contributed_amount",
        "remaining_amount",
        "progress_pct",
    ),
    "goal_contributions": ("goal_id", "month", "amount"),
    "paycheck_plan": (
        "paycheck_amount",
        "allocation_mode",
        "fixed_obligations_total",
        "needs",
        "goals_amount",
        "discretionary",
        "safety_buffer_reserved",
        "minimum_emergency_buffer",
        "saved_at",
    ),
    "paycheck_allocations": (
        "goal_id",
        "name",
        "category",
        "priority",
        "target_date",
        "recommended_amount",
        "remaining_after_allocation",
        "required_per_paycheck",
        "feasible",
    ),
    "action_feedback": (
        "action_id",
        "status",
        "action_type",
        "updated_at",
        "cooldown_until",
        "snooze_until",
    ),
    "action_history": ("action_id", "action_type", "outcome", "at"),
    "action_type_bias": ("action_type", "bias"),
    "action_metric_weights": ("metric", "weight"),
    "action_catalog": (
        "action_id",
        "action_type",
        "title",
        "impact_score",
        "urgency_score",
        "confidence_score",
        "effort_score",
    ),
}


def test_export_column_names_stay_stable():
    assert tuple(DATASETS) == tuple(EXPECTED_COLUMNS)
    for name, columns in EXPECTED_COLUMNS.items():
        assert tuple(column.name for column in DATASETS[name]) == columns


def _canonicalize(dataset: str, rows: list[dict]) -> list[dict]:
    canonical = []
    for row in rows:
        item = {}
        for column in DATASETS[dataset]:
            value = row[column.name]
            if value is None or value == "":
                item[column.name] = None
            elif column.kind == "integer":
                item[column.name] = int(value)
            elif column.kind in {"money", "number"}:
                item[column.name] = float(value)
            elif column.kind == "boolean":
                if isinstance(value, str):
                    item[column.name] = value == "true"
                else:
                    item[column.name] = bool(value)
            else:
                item[column.name] = value
        canonical.append(item)
    return canonical


def _datasets_from_json(payload: dict) -> dict[str, list[dict]]:
    return {name: _canonicalize(name, payload[name]) for name in DATASET_NAMES}


def _datasets_from_csv_zip(content: bytes) -> tuple[dict, dict[str, list[dict]]]:
    import zipfile

    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        names = set(archive.namelist())
        assert names == {f"{name}.csv" for name in DATASET_NAMES} | {"backup.json"}
        backup = json.loads(archive.read("backup.json"))
        tables = {}
        for name in DATASET_NAMES:
            text = archive.read(f"{name}.csv").decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text))
            assert reader.fieldnames == list(EXPECTED_COLUMNS[name])
            tables[name] = _canonicalize(name, list(reader))
    return backup, tables


def _datasets_from_xlsx(content: bytes) -> dict[str, list[dict]]:
    workbook = load_workbook(io.BytesIO(content))
    assert workbook.sheetnames == list(DATASET_NAMES)
    tables = {}
    for name in DATASET_NAMES:
        sheet = workbook[name]
        rows = list(sheet.iter_rows(values_only=True))
        assert rows[0] == EXPECTED_COLUMNS[name]
        columns = list(EXPECTED_COLUMNS[name])
        records = [dict(zip(columns, row, strict=True)) for row in rows[1:]]
        tables[name] = _canonicalize(name, records)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    assert cell.data_type == "s"
    return tables


def _find(rows: list[dict], **expected) -> dict:
    for row in rows:
        if all(row[key] == value for key, value in expected.items()):
            return row
    raise AssertionError(f"No row matching {expected} in {rows}")


def _write_rules(path: Path, rules: list[tuple[str, str]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["category", "keywords"])
        writer.writerows(rules)


def _isolate_data_dirs(tmp_path, monkeypatch):
    """Point budget and category IO at a temp tree and snapshot the real files."""
    import sdk.data_paths as data_paths

    real_data = data_paths.DATA_DIR
    real_local = data_paths.LOCAL_DIR
    snapshots: dict[Path, bytes | None] = {}
    for path in (
        real_data / "budget.csv",
        real_data / "categories.csv",
        real_local / "budget.csv",
        real_local / "categories.csv",
    ):
        snapshots[path] = path.read_bytes() if path.exists() else None

    data = tmp_path / "data"
    local = data / "local"
    data.mkdir()
    local.mkdir()
    monkeypatch.setattr(data_paths, "DATA_DIR", data)
    monkeypatch.setattr(data_paths, "LOCAL_DIR", local)
    return data, local, snapshots


def _assert_real_data_unchanged(snapshots: dict[Path, bytes | None]) -> None:
    for path, original in snapshots.items():
        if original is None:
            assert not path.exists()
        else:
            assert path.read_bytes() == original


def test_full_export_round_trips_every_format(tmp_path, monkeypatch):
    data, local, snapshots = _isolate_data_dirs(tmp_path, monkeypatch)
    shipped_budget = "category,monthly_budget\nGroceries,10\n"
    shipped_rules = "category,keywords\nGroceries,walmart\n"
    (data / "budget.csv").write_text(shipped_budget, encoding="utf-8")
    (data / "categories.csv").write_text(shipped_rules, encoding="utf-8")
    _write_rules(
        local / "categories.csv",
        [
            ("Income", "payroll|salary"),
            ("Meals", "cafe|coffee"),
            ("Savings", "emergency"),
            ("Subscriptions", "streamflix"),
            ("Quoted", "alpha, beta"),
        ],
    )

    client = TestClient(app)
    csv_body = (
        "Date,Description,Amount\n"
        "2026-01-05,STREAMFLIX,-12.00\n"
        "2026-02-05,STREAMFLIX,-12.00\n"
        "2026-03-02,Corner Cafe,-4.50\n"
        "2026-03-05,STREAMFLIX,-12.00\n"
        "2026-03-18,Payroll Deposit,2500.00\n"
        "2026-03-20,Emergency Fund Transfer,-100.00\n"
        "2026-04-02,Café,-6.25\n"
        "2026-04-05,STREAMFLIX,-12.00\n"
        "2026-05-05,STREAMFLIX,-12.00\n"
        "2026-05-11,=2+2,-2.00\n"
        "2026-06-05,STREAMFLIX,-12.00\n"
    ).encode()
    upload = client.post("/api/upload", files=[("files", ("synthetic.csv", csv_body, "text/csv"))])
    assert upload.status_code == 200

    cafe = _find(
        client.get("/api/ledger").json()["transactions"],
        description="Corner Cafe",
    )
    edited = client.patch(
        f"/api/ledger/transactions/{cafe['id']}/category",
        json={"category": "Dining Out"},
    )
    assert edited.status_code == 200

    budget = client.put(
        "/api/budget",
        json={"budget": {"Meals": 80, "Subscriptions": 15.99, "Savings": 100}},
    )
    assert budget.status_code == 200

    created = client.post(
        "/api/goals",
        json={
            "name": "Emergency Fund",
            "target_amount": 1000,
            "target_date": "2026-12-31",
            "priority": 1,
            "category": "emergency",
            "status": "active",
        },
    )
    assert created.status_code == 200
    goal = created.json()["goal"]
    assert goal["contributed_amount"] == 100

    recommendation = client.post(
        "/api/goals/paycheck-plan",
        json={
            "paycheck_amount": 1800,
            "fixed_obligations": [{"name": "Rent", "amount": 900}],
            "safety_buffer": 100,
            "minimum_emergency_buffer": 0,
            "mode": "balanced",
            "paychecks_per_month": 2,
            "goal_ids": [goal["id"]],
        },
    )
    assert recommendation.status_code == 200
    recommended = recommendation.json()
    saved = client.post(
        "/api/goals/paycheck-plan/save",
        json={
            "paycheck_amount": recommended["paycheck_amount"],
            "fixed_obligations": [{"name": "Rent", "amount": 900}],
            "safety_buffer_reserved": recommended["safety_buffer_reserved"],
            "minimum_emergency_buffer": 0,
            "mode": recommended["allocation_mode"],
            "needs": recommended["needs"],
            "goals": recommended["goals"],
            "discretionary": recommended["discretionary"],
            "goal_allocations": recommended["goal_allocations"],
        },
    )
    assert saved.status_code == 200

    listing = client.get("/api/subscriptions").json()["subscriptions"]
    stream = next(item for item in listing if item["merchant"] == "STREAMFLIX")
    preference = client.post(
        f"/api/subscriptions/{stream['stream_id']}/preferences",
        json={"essential": True},
    )
    assert preference.status_code == 200

    session_id = client.cookies.get("session_id")
    assert session_id
    action_state = get_action_state(session_id)
    action_state.clear()
    action_state.update(
        {
            "metric_weights": {
                "impact": 0.4,
                "urgency": 0.3,
                "confidence": 0.2,
                "effort": 0.1,
            },
            "type_bias": {"subscription_cleanup": -0.04},
            "history": [
                {
                    "action_id": "abc123",
                    "action_type": "subscription_cleanup",
                    "outcome": "dismissed",
                    "at": "2026-03-01T12:00:00+00:00",
                }
            ],
            "states": {
                "abc123": {
                    "status": "dismissed",
                    "updated_at": "2026-03-01T12:00:00+00:00",
                    "action_type": "subscription_cleanup",
                    "cooldown_until": "2026-03-04",
                }
            },
            "catalog": {
                "abc123": {
                    "action_type": "subscription_cleanup",
                    "title": "Review STREAMFLIX",
                    "impact_score": 0.5,
                    "urgency_score": 0.4,
                    "confidence_score": 0.8,
                    "effort_score": 0.2,
                }
            },
        }
    )

    exported_on = datetime.now(UTC).date().isoformat()
    json_response = client.get("/api/export?format=json")
    xlsx_response = client.get("/api/export?format=xlsx")
    zip_response = client.get("/api/export?format=zip")
    assert json_response.status_code == 200
    assert xlsx_response.status_code == 200
    assert zip_response.status_code == 200
    assert json_response.headers["content-disposition"] == (
        f'attachment; filename="finance-export-{exported_on}.json"'
    )
    assert xlsx_response.headers["content-disposition"] == (
        f'attachment; filename="finance-export-{exported_on}.xlsx"'
    )
    assert zip_response.headers["content-disposition"] == (
        f'attachment; filename="finance-export-{exported_on}.zip"'
    )

    document = json.loads(json_response.content)
    json_tables = _datasets_from_json(document)
    zip_document, zip_tables = _datasets_from_csv_zip(zip_response.content)
    xlsx_tables = _datasets_from_xlsx(xlsx_response.content)
    assert _datasets_from_json(zip_document) == zip_tables
    assert json_tables == zip_tables
    assert json_tables == xlsx_tables

    cafe_row = _find(json_tables["transactions"], description="Corner Cafe")
    assert cafe_row["date"] == "2026-03-02"
    assert cafe_row["amount"] == -4.5
    assert isinstance(cafe_row["amount"], float)
    assert cafe_row["category"] == "Dining Out"
    assert cafe_row["category_edited"] is True
    assert cafe_row["source_file"] == "synthetic.csv"
    formula = _find(json_tables["transactions"], description="=2+2")
    assert formula["amount"] == -2
    assert _find(json_tables["transactions"], description="Café")["amount"] == -6.25
    assert _find(json_tables["budgets"], category="Subscriptions")["monthly_budget"] == 15.99
    assert _find(json_tables["category_rules"], category="Quoted")["keywords"] == "alpha, beta"
    subscription = _find(json_tables["subscriptions"], merchant="STREAMFLIX")
    assert subscription["amount"] == 12
    assert subscription["essential"] is True
    assert subscription["ignored"] is False
    assert subscription["dominant_category"] == "Subscriptions"
    charges = [
        row
        for row in json_tables["subscription_charges"]
        if row["stream_id"] == subscription["stream_id"]
    ]
    assert [row["date"] for row in charges] == [
        "2026-01-05",
        "2026-02-05",
        "2026-03-05",
        "2026-04-05",
        "2026-05-05",
        "2026-06-05",
    ]
    assert all(row["amount"] == 12 for row in charges)
    goal_row = _find(json_tables["goals"], name="Emergency Fund")
    assert goal_row["target_amount"] == 1000
    assert goal_row["contributed_amount"] == 100
    assert goal_row["remaining_amount"] == 900
    assert goal_row["target_date"] == "2026-12-31"
    assert _find(json_tables["goal_contributions"], goal_id=goal_row["id"])["month"] == "2026-03"
    assert json_tables["paycheck_plan"][0]["paycheck_amount"] == 1800
    assert json_tables["paycheck_plan"][0]["goals_amount"] == saved.json()["plan"]["goals"]
    assert json_tables["paycheck_allocations"][0]["goal_id"] == goal["id"]
    feedback = _find(json_tables["action_feedback"], action_id="abc123")
    assert feedback["status"] == "dismissed"
    assert feedback["cooldown_until"] == "2026-03-04"
    assert feedback["snooze_until"] is None
    assert _find(json_tables["action_catalog"], action_id="abc123")["title"] == "Review STREAMFLIX"
    assert "$" not in json_response.text

    preview = client.post("/api/export/preview", json=document)
    assert preview.status_code == 200
    assert preview.json()["counts"]["transactions"] == len(json_tables["transactions"])
    assert preview.json()["warnings"]

    mutated_budget = client.put("/api/budget", json={"budget": {"Meals": 1}})
    assert mutated_budget.status_code == 200
    client.put(
        f"/api/goals/{goal['id']}",
        json={
            "name": "Renamed Goal",
            "target_amount": 5,
            "target_date": "2026-12-31",
            "priority": 2,
            "category": "emergency",
            "status": "paused",
        },
    )
    replacement = b"Date,Description,Amount\n2026-01-01,Other,-1.00\n"
    client.post("/api/upload", files=[("files", ("other.csv", replacement, "text/csv"))])
    _write_rules(local / "categories.csv", [("Groceries", "market")])
    get_action_state(session_id).clear()
    mutated = client.get("/api/budget").json()["budget"]
    assert mutated == [{"category": "Meals", "monthly_budget": 1.0}]
    assert client.get("/api/goals").json()["goals"][0]["name"] == "Renamed Goal"

    untouched = client.post("/api/export/preview", json=document)
    assert untouched.status_code == 200
    assert client.get("/api/budget").json()["budget"] == mutated

    restored = client.post("/api/export/restore", json=document)
    assert restored.status_code == 200
    assert restored.json()["status"] == "restored"
    again = json.loads(client.get("/api/export?format=json").content)
    assert _datasets_from_json(again) == json_tables
    assert client.get("/api/budget").json()["budget"] == [
        {"category": "Meals", "monthly_budget": 80.0},
        {"category": "Subscriptions", "monthly_budget": 15.99},
        {"category": "Savings", "monthly_budget": 100.0},
    ]
    restored_ledger = client.get("/api/ledger").json()["transactions"]
    assert _find(restored_ledger, description="Corner Cafe")["category"] == "Dining Out"
    assert client.get("/api/goals").json()["goals"][0]["name"] == "Emergency Fund"
    restored_subs = client.get("/api/subscriptions").json()["subscriptions"]
    assert next(item["essential"] for item in restored_subs if item["merchant"] == "STREAMFLIX")
    assert client.get("/api/categories").json()["categories"][-1] == {
        "category": "Quoted",
        "keywords": "alpha, beta",
    }
    assert (data / "budget.csv").read_text(encoding="utf-8") == shipped_budget
    assert (data / "categories.csv").read_text(encoding="utf-8") == shipped_rules
    assert "Meals" in (local / "budget.csv").read_text(encoding="utf-8")
    assert "alpha, beta" in (local / "categories.csv").read_text(encoding="utf-8")
    _assert_real_data_unchanged(snapshots)


def test_transaction_export_filters_by_date_and_category(tmp_path, monkeypatch):
    data, _local, snapshots = _isolate_data_dirs(tmp_path, monkeypatch)
    _write_rules(
        data / "categories.csv",
        [
            ("Meals", "cafe"),
            ("Income", "payroll"),
            ("Shopping", "book"),
        ],
    )
    client = TestClient(app)
    csv_body = (
        b"Date,Description,Amount\n"
        b"2026-01-15,Cafe,-3.00\n"
        b"2026-02-15,Cafe,-4.00\n"
        b"2026-02-16,Payroll,100.00\n"
        b"2026-03-02,Mystery Kiosk,-9.00\n"
    )
    uploaded = client.post("/api/upload", files=[("files", ("filter.csv", csv_body, "text/csv"))])
    assert uploaded.status_code == 200

    ranged = client.get(
        "/api/ledger/transactions/export",
        params={
            "all_dates": "true",
            "start_date": "2026-02-01",
            "end_date": "2026-02-28",
            "categories": ["Meals", "Shopping"],
        },
    )
    assert ranged.status_code == 200
    exported_on = datetime.now(UTC).date().isoformat()
    assert (
        ranged.headers["content-disposition"]
        == f'attachment; filename="transactions-2026-02-01-to-2026-02-28-{exported_on}.csv"'
    )
    assert "$" not in ranged.text
    assert "Cafe" in ranged.text
    assert "Payroll" not in ranged.text
    assert "Mystery Kiosk" not in ranged.text
    reader = csv.DictReader(io.StringIO(ranged.text))
    assert reader.fieldnames == ["id", "date", "description", "amount", "category", "source_file"]
    amounts = [float(row["amount"]) for row in reader]
    assert amounts == [-4.0]

    meals = client.get(
        "/api/ledger/transactions/export",
        params={"all_dates": "true", "category": "Meals", "format": "xlsx"},
    )
    assert meals.status_code == 200
    workbook = load_workbook(io.BytesIO(meals.content))
    values = list(workbook["Transactions"].iter_rows(values_only=True))
    assert values[0] == ("id", "date", "description", "amount", "category", "source_file")
    assert [row[2] for row in values[1:]] == ["Cafe", "Cafe"]
    amounts = [row[3] for row in values[1:]]
    assert amounts
    assert all(
        isinstance(amount, int | float) and not isinstance(amount, bool) for amount in amounts
    )

    uncategorized = client.get(
        "/api/ledger/transactions",
        params={"all_dates": "true", "categories": ["Uncategorized"]},
    )
    assert uncategorized.status_code == 200
    descriptions = [row["description"] for row in uncategorized.json()["transactions"]]
    assert descriptions == ["Mystery Kiosk"]

    latest = client.get("/api/ledger/transactions/export")
    assert latest.status_code == 200
    assert f"transactions-2026-03-{exported_on}.csv" in latest.headers["content-disposition"]
    assert "Mystery Kiosk" in latest.text
    assert "Cafe" not in latest.text

    invalid = client.get(
        "/api/ledger/transactions/export",
        params={"start_date": "2026-03-02", "end_date": "2026-03-01"},
    )
    assert invalid.status_code == 422
    bad_date = client.get("/api/ledger/transactions/export", params={"start_date": "2026-02-31"})
    assert bad_date.status_code == 422
    _assert_real_data_unchanged(snapshots)


def test_backup_preview_rejects_invalid_documents():
    client = TestClient(app)
    response = client.post("/api/export/preview", json={"version": 1})
    assert response.status_code == 422
    with pytest.raises(BackupValidationError):
        parse_backup({"version": 2})
