"""Full-account export and JSON backup restore.

Datasets use stable column names. Dates are ISO strings (``YYYY-MM-DD`` or
``YYYY-MM``). Timestamps are ISO-8601. Amounts are numbers, never currency
text. CSV files are UTF-8 with a byte-order mark. The JSON backup is the
lossless copy; spreadsheet files carry the same rows.

Subscription charge rows and goal contribution rows are included so a
spreadsheet is readable on its own. Restore keeps user-edited records and
rebuilds those derived rows from transactions and goals.

Paycheck obligation names are not stored by the app, so the paycheck plan
sheet keeps the obligation total only.
"""

from __future__ import annotations

import csv
import io
import json
import math
import re
import zipfile
from collections.abc import Mapping
from dataclasses import dataclass
from datetime import datetime
from typing import Literal

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

BACKUP_VERSION = 1
MAX_ROWS = 100_000
MAX_KEYWORD_LENGTH = 4_000

Kind = Literal[
    "string",
    "money",
    "number",
    "integer",
    "boolean",
    "date",
    "datetime",
    "month",
]


@dataclass(frozen=True)
class Column:
    name: str
    kind: Kind
    nullable: bool = False


def _columns(*items: tuple[str, Kind] | tuple[str, Kind, bool]) -> tuple[Column, ...]:
    columns: list[Column] = []
    for item in items:
        if len(item) == 2:
            name, kind = item
            columns.append(Column(name, kind))
        else:
            name, kind, nullable = item
            columns.append(Column(name, kind, nullable))
    return tuple(columns)


DATASETS: dict[str, tuple[Column, ...]] = {
    "transactions": _columns(
        ("id", "string"),
        ("date", "date"),
        ("description", "string"),
        ("amount", "money"),
        ("category", "string"),
        ("source_file", "string"),
        ("category_edited", "boolean"),
    ),
    "budgets": _columns(
        ("category", "string"),
        ("monthly_budget", "money"),
    ),
    "category_rules": _columns(
        ("category", "string"),
        ("keywords", "string"),
    ),
    "subscriptions": _columns(
        ("stream_id", "string"),
        ("merchant", "string"),
        ("dominant_category", "string"),
        ("cadence", "string"),
        ("amount", "money"),
        ("baseline_amount", "money"),
        ("expected_amount", "money"),
        ("last_paid_amount", "money", True),
        ("confidence", "number"),
        ("active", "boolean"),
        ("ignored", "boolean"),
        ("essential", "boolean"),
        ("manually_managed", "boolean"),
        ("price_increase", "boolean"),
        ("cancellation_candidate", "boolean"),
        ("negotiation_opportunity", "boolean"),
        ("is_new_recurring", "boolean"),
        ("missed_expected_charge", "boolean"),
        ("charge_count", "integer"),
        ("trend", "string"),
        ("last_charge_date", "date"),
        ("next_expected_charge_date", "date", True),
        ("next_due_date", "date", True),
        ("status_group", "string", True),
        ("payment_state", "string", True),
    ),
    "subscription_charges": _columns(
        ("stream_id", "string"),
        ("date", "date"),
        ("amount", "money"),
    ),
    "goals": _columns(
        ("id", "string"),
        ("name", "string"),
        ("target_amount", "money"),
        ("target_date", "date", True),
        ("priority", "integer"),
        ("category", "string"),
        ("status", "string"),
        ("created_at", "datetime"),
        ("updated_at", "datetime"),
        ("contributed_amount", "money"),
        ("remaining_amount", "money"),
        ("progress_pct", "number"),
    ),
    "goal_contributions": _columns(
        ("goal_id", "string"),
        ("month", "month"),
        ("amount", "money"),
    ),
    "paycheck_plan": _columns(
        ("paycheck_amount", "money"),
        ("allocation_mode", "string"),
        ("fixed_obligations_total", "money"),
        ("needs", "money"),
        ("goals_amount", "money"),
        ("discretionary", "money"),
        ("safety_buffer_reserved", "money"),
        ("minimum_emergency_buffer", "money"),
        ("saved_at", "datetime"),
    ),
    "paycheck_allocations": _columns(
        ("goal_id", "string"),
        ("name", "string"),
        ("category", "string"),
        ("priority", "integer"),
        ("target_date", "date", True),
        ("recommended_amount", "money"),
        ("remaining_after_allocation", "money"),
        ("required_per_paycheck", "money", True),
        ("feasible", "boolean", True),
    ),
    "action_feedback": _columns(
        ("action_id", "string"),
        ("status", "string"),
        ("action_type", "string"),
        ("updated_at", "datetime"),
        ("cooldown_until", "date", True),
        ("snooze_until", "date", True),
    ),
    "action_history": _columns(
        ("action_id", "string"),
        ("action_type", "string"),
        ("outcome", "string"),
        ("at", "datetime"),
    ),
    "action_type_bias": _columns(
        ("action_type", "string"),
        ("bias", "number"),
    ),
    "action_metric_weights": _columns(
        ("metric", "string"),
        ("weight", "number"),
    ),
    "action_catalog": _columns(
        ("action_id", "string"),
        ("action_type", "string"),
        ("title", "string"),
        ("impact_score", "number"),
        ("urgency_score", "number"),
        ("confidence_score", "number"),
        ("effort_score", "number"),
    ),
}

DATASET_NAMES: tuple[str, ...] = tuple(DATASETS)
DOCUMENT_KEYS = frozenset({"version", "exported_at", *DATASET_NAMES})

RESTORE_WARNINGS: tuple[str, ...] = (
    "Restore replaces transactions, goals, the saved paycheck plan, "
    "subscription preferences, and action feedback in this session.",
    "Restore replaces the monthly budgets and category keyword rules saved for this app.",
    "Subscription lists, subscription charges, and goal progress are rebuilt "
    "from transactions and goals. Essential and ignored flags are restored.",
)


class BackupValidationError(ValueError):
    """The backup JSON is not a document this app can restore."""

    def __init__(self, errors: list[str]):
        self.errors = errors
        summary = "; ".join(errors[:8])
        remaining = len(errors) - 8
        if remaining > 0:
            summary = f"{summary}; and {remaining} more"
        super().__init__(summary or "Backup is not valid.")


def column_names(dataset: str) -> list[str]:
    return [column.name for column in DATASETS[dataset]]


def _is_missing(value: object) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return False
    if isinstance(value, (list, dict, tuple, set)):
        return False
    try:
        missing = pd.isna(value)
    except (TypeError, ValueError):
        return False
    return bool(missing)


def _is_bool(value: object) -> bool:
    if isinstance(value, bool):
        return True
    return type(value).__name__ in {"bool_", "bool"}


def _check_date(value: object, pattern: str) -> str:
    if not isinstance(value, str):
        raise ValueError("expected a string")
    try:
        parsed = datetime.strptime(value, pattern)
    except ValueError as exc:
        raise ValueError(f"expected {pattern}") from exc
    if parsed.strftime(pattern) != value:
        raise ValueError(f"expected {pattern}")
    return value


def _check_timestamp(value: object) -> str:
    if not isinstance(value, str) or "T" not in value:
        raise ValueError("expected an ISO-8601 timestamp")
    text = f"{value[:-1]}+00:00" if value.endswith("Z") else value
    try:
        datetime.fromisoformat(text)
    except ValueError as exc:
        raise ValueError("expected an ISO-8601 timestamp") from exc
    return value


def _finite_number(value: object) -> float:
    if _is_bool(value) or isinstance(value, str):
        raise ValueError("expected a number")
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError("expected a number") from exc
    if not math.isfinite(number):
        raise ValueError("expected a finite number")
    return number


def coerce_value(column: Column, value: object) -> object:
    if _is_missing(value):
        if column.nullable:
            return None
        raise ValueError("is required")
    if column.kind == "boolean":
        if not _is_bool(value):
            raise ValueError("expected true or false")
        return bool(value)
    if column.kind == "integer":
        number = _finite_number(value)
        if not number.is_integer():
            raise ValueError("expected an integer")
        return int(number)
    if column.kind == "money":
        return round(_finite_number(value), 2)
    if column.kind == "number":
        return round(_finite_number(value), 6)
    if column.kind == "string":
        if not isinstance(value, str):
            raise ValueError("expected a string")
        return value
    if column.kind == "date":
        return _check_date(value, "%Y-%m-%d")
    if column.kind == "month":
        return _check_date(value, "%Y-%m")
    if column.kind == "datetime":
        return _check_timestamp(value)
    raise ValueError("unknown column type")


def _coerce_row(dataset: str, row: object, index: int, errors: list[str]) -> dict | None:
    if not isinstance(row, Mapping):
        errors.append(f"{dataset}[{index}] must be an object")
        return None
    cleaned: dict[str, object] = {}
    for column in DATASETS[dataset]:
        if column.name not in row:
            errors.append(f"{dataset}[{index}].{column.name} is required")
            continue
        try:
            cleaned[column.name] = coerce_value(column, row[column.name])
        except ValueError as exc:
            errors.append(f"{dataset}[{index}].{column.name} {exc}")
    return cleaned


def _require_unique(
    dataset: str,
    rows: list[dict],
    fields: tuple[str, ...],
    errors: list[str],
) -> None:
    seen: set[tuple[object, ...]] = set()
    for index, row in enumerate(rows):
        key = tuple(row.get(field) for field in fields)
        if key in seen:
            label = ", ".join(fields)
            rendered = key[0] if len(key) == 1 else key
            errors.append(f"{dataset}[{index}] duplicates {label} {rendered}")
            continue
        seen.add(key)


def _validate_keywords(rules: list[dict], errors: list[str]) -> None:
    for index, rule in enumerate(rules):
        keywords = str(rule.get("keywords", ""))
        if len(keywords) > MAX_KEYWORD_LENGTH:
            errors.append(f"category_rules[{index}].keywords is too long")
            continue
        if not keywords.strip():
            continue
        try:
            re.compile(keywords)
        except re.error:
            errors.append(f"category_rules[{index}].keywords is not a valid pattern")


def parse_backup(payload: object) -> dict:
    """Validate a backup document and return canonical rows."""
    errors: list[str] = []
    if not isinstance(payload, Mapping):
        raise BackupValidationError(["Backup must be a JSON object."])

    extra = sorted(set(payload) - DOCUMENT_KEYS)
    missing = sorted(DOCUMENT_KEYS - set(payload))
    if extra:
        errors.append(f"Unexpected fields: {', '.join(extra)}")
    if missing:
        errors.append(f"Missing fields: {', '.join(missing)}")
    if payload.get("version") != BACKUP_VERSION:
        errors.append("version must be 1")
    try:
        exported_at = _check_timestamp(payload.get("exported_at"))
    except ValueError:
        exported_at = ""
        errors.append("exported_at must be an ISO-8601 timestamp")

    datasets: dict[str, list[dict]] = {}
    for name in DATASET_NAMES:
        raw_rows = payload.get(name)
        if not isinstance(raw_rows, list):
            if name in payload:
                errors.append(f"{name} must be a list")
            datasets[name] = []
            continue
        if len(raw_rows) > MAX_ROWS:
            errors.append(f"{name} exceeds {MAX_ROWS} rows")
            datasets[name] = []
            continue
        cleaned_rows: list[dict] = []
        for index, raw_row in enumerate(raw_rows):
            cleaned = _coerce_row(name, raw_row, index, errors)
            if cleaned is not None and len(cleaned) == len(DATASETS[name]):
                cleaned_rows.append(cleaned)
        datasets[name] = cleaned_rows

    _require_unique("transactions", datasets["transactions"], ("id",), errors)
    _require_unique("budgets", datasets["budgets"], ("category",), errors)
    _require_unique("subscriptions", datasets["subscriptions"], ("stream_id",), errors)
    _require_unique("goals", datasets["goals"], ("id",), errors)
    _require_unique(
        "goal_contributions",
        datasets["goal_contributions"],
        ("goal_id", "month"),
        errors,
    )
    _require_unique("action_feedback", datasets["action_feedback"], ("action_id",), errors)
    _require_unique("action_type_bias", datasets["action_type_bias"], ("action_type",), errors)
    _require_unique("action_metric_weights", datasets["action_metric_weights"], ("metric",), errors)
    _require_unique("action_catalog", datasets["action_catalog"], ("action_id",), errors)
    _validate_keywords(datasets["category_rules"], errors)

    if len(datasets["paycheck_plan"]) > 1:
        errors.append("paycheck_plan must contain at most one row")
    if datasets["paycheck_allocations"] and not datasets["paycheck_plan"]:
        errors.append("paycheck_allocations require a paycheck_plan row")

    if errors:
        raise BackupValidationError(errors)

    return {
        "version": BACKUP_VERSION,
        "exported_at": exported_at,
        **datasets,
    }


def _row_from_mapping(dataset: str, source: Mapping) -> dict:
    errors: list[str] = []
    cleaned = _coerce_row(dataset, source, 0, errors)
    if cleaned is None or errors:
        raise BackupValidationError(errors or [f"{dataset} row is incomplete"])
    return cleaned


def assemble_backup(
    *,
    exported_at: str,
    transactions: list[Mapping],
    budgets: list[Mapping],
    category_rules: list[Mapping],
    subscriptions: list[Mapping],
    goals: list[Mapping],
    paycheck_plan: Mapping | None,
    action_state: Mapping | None,
) -> dict:
    """Build a canonical backup from in-app records."""
    subscription_rows: list[dict] = []
    charge_rows: list[dict] = []
    for subscription in subscriptions:
        history = subscription.get("charge_history") or []
        subscription_rows.append(_row_from_mapping("subscriptions", subscription))
        if isinstance(history, list):
            for charge in history:
                if isinstance(charge, Mapping):
                    charge_rows.append(
                        _row_from_mapping(
                            "subscription_charges",
                            {
                                "stream_id": subscription.get("stream_id", ""),
                                "date": charge.get("date"),
                                "amount": charge.get("amount"),
                            },
                        )
                    )

    goal_rows: list[dict] = []
    contribution_rows: list[dict] = []
    for goal in goals:
        goal_rows.append(_row_from_mapping("goals", goal))
        history = goal.get("contribution_history") or []
        if isinstance(history, list):
            for point in history:
                if isinstance(point, Mapping):
                    contribution_rows.append(
                        _row_from_mapping(
                            "goal_contributions",
                            {
                                "goal_id": goal.get("id", ""),
                                "month": point.get("month"),
                                "amount": point.get("amount"),
                            },
                        )
                    )

    plan_rows: list[dict] = []
    allocation_rows: list[dict] = []
    if paycheck_plan:
        plan_source = dict(paycheck_plan)
        plan_source["goals_amount"] = paycheck_plan.get("goals_amount", paycheck_plan.get("goals"))
        plan_rows.append(_row_from_mapping("paycheck_plan", plan_source))
        allocations = paycheck_plan.get("goal_allocations") or []
        if isinstance(allocations, list):
            for allocation in allocations:
                if isinstance(allocation, Mapping):
                    allocation_rows.append(_row_from_mapping("paycheck_allocations", allocation))

    state = dict(action_state or {})
    feedback_rows = []
    states = state.get("states") or {}
    if isinstance(states, Mapping):
        for action_id in sorted(states):
            entry = states[action_id]
            if isinstance(entry, Mapping):
                feedback_rows.append(
                    _row_from_mapping(
                        "action_feedback",
                        {
                            "action_id": str(action_id),
                            "status": entry.get("status"),
                            "action_type": entry.get("action_type"),
                            "updated_at": entry.get("updated_at"),
                            "cooldown_until": entry.get("cooldown_until"),
                            "snooze_until": entry.get("snooze_until"),
                        },
                    )
                )

    history_rows = []
    history = state.get("history") or []
    if isinstance(history, list):
        for entry in history:
            if isinstance(entry, Mapping):
                history_rows.append(_row_from_mapping("action_history", entry))
    history_rows.sort(key=lambda row: (str(row["at"]), str(row["action_id"]), str(row["outcome"])))

    bias_rows = []
    bias = state.get("type_bias") or {}
    if isinstance(bias, Mapping):
        for action_type in sorted(bias):
            bias_rows.append(
                _row_from_mapping(
                    "action_type_bias",
                    {"action_type": str(action_type), "bias": bias[action_type]},
                )
            )

    weight_rows = []
    weights = state.get("metric_weights") or {}
    if isinstance(weights, Mapping):
        for metric in sorted(weights):
            weight_rows.append(
                _row_from_mapping(
                    "action_metric_weights",
                    {"metric": str(metric), "weight": weights[metric]},
                )
            )

    catalog_rows = []
    catalog = state.get("catalog") or {}
    if isinstance(catalog, Mapping):
        for action_id in sorted(catalog):
            entry = catalog[action_id]
            if isinstance(entry, Mapping):
                catalog_rows.append(
                    _row_from_mapping(
                        "action_catalog",
                        {"action_id": str(action_id), **dict(entry)},
                    )
                )

    document = {
        "version": BACKUP_VERSION,
        "exported_at": _check_timestamp(exported_at),
        "transactions": [_row_from_mapping("transactions", row) for row in transactions],
        "budgets": [_row_from_mapping("budgets", row) for row in budgets],
        "category_rules": [_row_from_mapping("category_rules", row) for row in category_rules],
        "subscriptions": subscription_rows,
        "subscription_charges": charge_rows,
        "goals": goal_rows,
        "goal_contributions": contribution_rows,
        "paycheck_plan": plan_rows,
        "paycheck_allocations": allocation_rows,
        "action_feedback": feedback_rows,
        "action_history": history_rows,
        "action_type_bias": bias_rows,
        "action_metric_weights": weight_rows,
        "action_catalog": catalog_rows,
    }
    return parse_backup(document)


def backup_counts(document: Mapping) -> dict[str, int]:
    return {name: len(document[name]) for name in DATASET_NAMES}


def backup_preview(document: Mapping) -> dict:
    return {
        "version": BACKUP_VERSION,
        "exported_at": document["exported_at"],
        "counts": backup_counts(document),
        "warnings": list(RESTORE_WARNINGS),
    }


def dumps_backup(document: Mapping) -> str:
    ordered = {"version": document["version"], "exported_at": document["exported_at"]}
    for name in DATASET_NAMES:
        ordered[name] = document[name]
    return json.dumps(ordered, indent=2, ensure_ascii=False, allow_nan=False) + "\n"


def _csv_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        text = f"{value:.6f}".rstrip("0").rstrip(".")
        return "0" if text in {"", "-"} else text
    return str(value)


def render_csv(dataset: str, rows: list[Mapping]) -> str:
    buffer = io.StringIO()
    writer = csv.DictWriter(
        buffer,
        fieldnames=column_names(dataset),
        lineterminator="\n",
        extrasaction="ignore",
    )
    writer.writeheader()
    for row in rows:
        writer.writerow({name: _csv_cell(row.get(name)) for name in column_names(dataset)})
    return buffer.getvalue()


def _write_cell(
    sheet: Worksheet,
    row_index: int,
    column_index: int,
    column: Column,
    value: object,
) -> None:
    cell: Cell = sheet.cell(row=row_index, column=column_index, value=None)
    if value is None:
        return
    if column.kind == "boolean":
        cell.value = bool(value)
        cell.data_type = "b"
        return
    if column.kind in {"money", "number", "integer"}:
        cell.value = value
        if column.kind == "money":
            cell.number_format = "0.00"
        elif column.kind == "integer":
            cell.number_format = "0"
        else:
            cell.number_format = "0.######"
        return
    cell.value = str(value)
    cell.data_type = "s"


def render_xlsx(document: Mapping) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError("Workbook has no active sheet.")
    for index, name in enumerate(DATASET_NAMES):
        current = sheet if index == 0 else workbook.create_sheet(name)
        if index == 0:
            current.title = name
        columns = DATASETS[name]
        for column_index, column in enumerate(columns, start=1):
            header = current.cell(row=1, column=column_index, value=column.name)
            header.data_type = "s"
            header.font = Font(bold=True)
        for row_index, row in enumerate(document[name], start=2):
            for column_index, column in enumerate(columns, start=1):
                _write_cell(current, row_index, column_index, column, row.get(column.name))
        current.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_zip(document: Mapping) -> bytes:
    payload = dumps_backup(document).encode("utf-8")
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name in DATASET_NAMES:
            csv_text = render_csv(name, document[name])
            archive.writestr(f"{name}.csv", csv_text.encode("utf-8-sig"))
        archive.writestr("backup.json", payload)
    return buffer.getvalue()


def materialize_transactions(document: Mapping) -> tuple[pd.DataFrame, dict[str, str]]:
    rows = document["transactions"]
    frame = pd.DataFrame(
        {
            "transaction_id": [row["id"] for row in rows],
            "date": pd.to_datetime([row["date"] for row in rows]),
            "description": [row["description"] for row in rows],
            "amount": [row["amount"] for row in rows],
            "category": [row["category"] for row in rows],
            "source_file": [row["source_file"] for row in rows],
        }
    )
    overrides = {row["id"]: row["category"] for row in rows if row["category_edited"]}
    return frame, overrides


def materialize_budgets(document: Mapping) -> dict[str, float]:
    return {row["category"]: float(row["monthly_budget"]) for row in document["budgets"]}


def materialize_category_rules(document: Mapping) -> list[dict[str, str]]:
    return [
        {"category": row["category"], "keywords": row["keywords"]}
        for row in document["category_rules"]
    ]


def materialize_subscription_preferences(document: Mapping) -> dict[str, dict[str, bool]]:
    return {
        row["stream_id"]: {"essential": bool(row["essential"]), "ignored": bool(row["ignored"])}
        for row in document["subscriptions"]
    }


def materialize_goals(document: Mapping) -> list[dict]:
    return [
        {
            "id": row["id"],
            "name": row["name"],
            "target_amount": float(row["target_amount"]),
            "target_date": row["target_date"],
            "priority": int(row["priority"]),
            "category": row["category"],
            "status": row["status"],
            "created_at": row["created_at"],
            "updated_at": row["updated_at"],
        }
        for row in document["goals"]
    ]


def materialize_paycheck_plan(document: Mapping) -> dict | None:
    if not document["paycheck_plan"]:
        return None
    plan = document["paycheck_plan"][0]
    return {
        "paycheck_amount": float(plan["paycheck_amount"]),
        "allocation_mode": plan["allocation_mode"],
        "fixed_obligations_total": float(plan["fixed_obligations_total"]),
        "needs": float(plan["needs"]),
        "goals": float(plan["goals_amount"]),
        "discretionary": float(plan["discretionary"]),
        "safety_buffer_reserved": float(plan["safety_buffer_reserved"]),
        "minimum_emergency_buffer": float(plan["minimum_emergency_buffer"]),
        "goal_allocations": [
            {
                "goal_id": row["goal_id"],
                "name": row["name"],
                "category": row["category"],
                "priority": int(row["priority"]),
                "target_date": row["target_date"],
                "recommended_amount": float(row["recommended_amount"]),
                "remaining_after_allocation": float(row["remaining_after_allocation"]),
                "required_per_paycheck": (
                    None
                    if row["required_per_paycheck"] is None
                    else float(row["required_per_paycheck"])
                ),
                "feasible": row["feasible"],
            }
            for row in document["paycheck_allocations"]
        ],
        "saved_at": plan["saved_at"],
    }


def materialize_action_state(document: Mapping) -> dict:
    tables = (
        "action_feedback",
        "action_history",
        "action_type_bias",
        "action_metric_weights",
        "action_catalog",
    )
    if not any(document[name] for name in tables):
        return {}
    states: dict[str, dict] = {}
    for row in document["action_feedback"]:
        entry = {
            "status": row["status"],
            "updated_at": row["updated_at"],
            "action_type": row["action_type"],
        }
        if row["cooldown_until"] is not None:
            entry["cooldown_until"] = row["cooldown_until"]
        if row["snooze_until"] is not None:
            entry["snooze_until"] = row["snooze_until"]
        states[row["action_id"]] = entry
    return {
        "metric_weights": {
            row["metric"]: float(row["weight"]) for row in document["action_metric_weights"]
        },
        "type_bias": {
            row["action_type"]: float(row["bias"]) for row in document["action_type_bias"]
        },
        "history": [
            {
                "action_id": row["action_id"],
                "action_type": row["action_type"],
                "outcome": row["outcome"],
                "at": row["at"],
            }
            for row in document["action_history"]
        ],
        "states": states,
        "catalog": {
            row["action_id"]: {
                "action_type": row["action_type"],
                "title": row["title"],
                "impact_score": float(row["impact_score"]),
                "urgency_score": float(row["urgency_score"]),
                "confidence_score": float(row["confidence_score"]),
                "effort_score": float(row["effort_score"]),
            }
            for row in document["action_catalog"]
        },
    }
